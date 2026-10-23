"""Erzeugt app/places.json: alle Städte, Märkte und Gemeinden in Oberfranken.
Quelle: Statistisches Bundesamt (Destatis), Gemeindeverzeichnis-Informationssystem GV-ISys,
„Gemeinden in Deutschland nach Fläche, Bevölkerung und Postleitzahl“ (AuszugGV2QAktuell.xlsx), Datenlizenz Deutschland – Namensnennung 2.0.
Aufruf: python tools/build_places.py <Pfad zur xlsx>"""
import json, os, re, sys
import openpyxl

KREIS = {"61": ("Bamberg", "kreisfreie Stadt Bamberg"), "62": ("Bayreuth", "kreisfreie Stadt Bayreuth"), "63": ("Coburg", "kreisfreie Stadt Coburg"),
         "64": ("Hof", "kreisfreie Stadt Hof"), "71": ("Bamberg", "Landkreis Bamberg"), "72": ("Bayreuth", "Landkreis Bayreuth"),
         "73": ("Coburg", "Landkreis Coburg"), "74": ("Forchheim", "Landkreis Forchheim"), "75": ("Hof", "Landkreis Hof"),
         "76": ("Kronach", "Landkreis Kronach"), "77": ("Kulmbach", "Landkreis Kulmbach"), "78": ("Lichtenfels", "Landkreis Lichtenfels"),
         "79": ("Wunsiedel", "Landkreis Wunsiedel i. Fichtelgebirge")}
ART = {"St": "Stadt", "M": "Markt", "GKSt": "Große Kreisstadt"}


def long_name(n):
    """Amtliche Abkürzungen ausschreiben: „Schwarzenbach a.d.Saale“ -> „Schwarzenbach an der Saale“."""
    n = re.sub(r"\bi\.\s*OFr\.?", "in Oberfranken", n)
    n = re.sub(r"\ba\.d\.\s*", "an der ", n)
    n = re.sub(r"\bi\.\s*", "im ", n)
    n = re.sub(r"\ba\.\s*", "am ", n)
    n = re.sub(r"\bb\.\s*", "bei ", n)
    n = re.sub(r"\bOFr\.?", "Oberfranken", n)
    return re.sub(r"\s+", " ", n).strip()


def main(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[[s for s in wb.sheetnames if s.startswith("Onlineprodukt")][0]]
    out = []
    for r in ws.iter_rows(values_only=True):
        if r[0] != "60" or r[2] != "09" or r[3] != "4" or not r[9]:
            continue  # nur Gemeinden in Oberfranken mit Einwohnern (ohne gemeindefreie Forstgebiete)
        raw = r[7]
        base, art = (raw.rsplit(",", 1) + [""])[:2]
        art = ART.get(art.strip(), "Gemeinde")
        kreis = r[4]
        region, kreis_name = KREIS[kreis]
        if kreis.startswith("6"):
            art = "kreisfreie Stadt"
        name = long_name(base.strip())
        out.append({"name": name, "amtlich": base.strip(), "art": art, "kreis": kreis_name, "region": region,
                    "lat": round(float(str(r[15]).replace(",", ".")), 5), "lon": round(float(str(r[14]).replace(",", ".")), 5),
                    "ew": int(r[9]), "plz": str(r[13] or ""), "ags": f"094{kreis}{r[6]}"})
    out.sort(key=lambda x: x["name"])
    meta = {"quelle": "Statistisches Bundesamt (Destatis), Gemeindeverzeichnis GV-ISys, Gebietsstand 30.06.2026, Bevölkerung 31.12.2024",
            "lizenz": "Datenlizenz Deutschland – Namensnennung – Version 2.0", "anzahl": len(out), "einwohner": sum(x["ew"] for x in out)}
    dst = os.path.join(os.path.dirname(__file__), "..", "app", "places.json")
    json.dump({"meta": meta, "orte": out}, open(dst, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(meta)


if __name__ == "__main__":
    main(sys.argv[1])
